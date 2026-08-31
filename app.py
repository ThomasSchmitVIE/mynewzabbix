from __future__ import annotations

import base64
import io
from zipfile import BadZipFile
import hashlib
import hmac
import json
import os
import threading
import time
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from bs4 import BeautifulSoup


APP_DIR = Path(__file__).resolve().parent
ZABBIX_URL = os.environ.get("ZABBIX_URL", "http://52.29.15.174/zabbix/api_jsonrpc.php")
CACHE_SECONDS = 45
SHEET_XLSX_URL = os.environ.get("LOCATION_SHEET_XLSX_URL", "https://docs.google.com/spreadsheets/d/1PzJVMVrupZ8aOGr0StOrtelEUYxdyMnk/export?format=xlsx")
cache_lock = threading.Lock()
problem_cache = {"time": 0.0, "data": None}
location_cache = {"time": 0.0, "locations": {}}
hidden_cache = {"time": 0.0, "problems": []}
sheet_cache = {"time": 0.0, "rows": {}}


class DashboardError(RuntimeError):
    pass


def zabbix_call(method, params, auth=None, request_id=1):
    payload = {"jsonrpc": "2.0", "method": method, "params": params, "id": request_id}
    if auth:
        payload["auth"] = auth
    response = requests.post(ZABBIX_URL, json=payload, timeout=25)
    response.raise_for_status()
    data = response.json()
    if "error" in data:
        raise DashboardError(data["error"].get("data") or data["error"].get("message"))
    return data["result"]


def category(description: str) -> str:
    text = description.lower()
    if "toner" in text:
        return "toner"
    if "jam" in text:
        return "paper_jam"
    if "out of paper" in text:
        return "paper_empty"
    if "low on paper" in text:
        return "paper_low"
    if "printer error" in text:
        return "printer_error"
    if "scanner" in text:
        return "scanner"
    if "uncompleted" in text:
        return "unfinished_print"
    if "unavailable" in text or "offline" in text or "icmp" in text:
        return "offline"
    return "other"


def clean_description(description: str, host: str) -> str:
    return description.replace("{HOST.NAME}", host).replace("{HOST.HOST}", host)


def printbox_login(session, config):
    response = session.get(config["login_url"], timeout=25); response.raise_for_status()
    tag = BeautifulSoup(response.text, "html.parser").find("meta", {"name": "csrf-token"})
    if not tag: raise DashboardError("Printbox-Anmeldung enthaelt kein Sicherheitstoken.")
    session.headers["X-CSRF-Token"] = tag["content"]
    response = session.post(config["login_url"], data=config["login_payload"], timeout=25); response.raise_for_status()


def classify_area(address):
    value = address.casefold()
    if "vienna" in value or "wien" in value: return "vienna"
    if "graz" in value: return "graz"
    if "linz" in value: return "linz"
    return "rest"


def location_fields(soup):
    fields = {}
    for row in soup.select("tr"):
        cells = row.find_all(["th", "td"])
        if len(cells) >= 2:
            fields[cells[0].get_text(" ", strip=True).casefold()] = cells[1].get_text(" ", strip=True)
    address = fields.get("location address", "")
    try:
        lat = float(fields.get("location lat", "").replace(",", "."))
        lon = float(fields.get("location long", "").replace(",", "."))
    except ValueError:
        lat, lon = None, None
    return {"area": classify_area(address), "address": address, "lat": lat, "lon": lon}


def fetch_locations(hosts):
    now = time.monotonic()
    with cache_lock:
        cached = location_cache["locations"]
        if cached and set(hosts).issubset(cached) and now - location_cache["time"] < 21600:
            return {host: cached[host] for host in hosts}
    try:
        config = json.loads(os.environ.get("PRINTBOX_CONFIG_JSON", ""))
        with requests.Session() as session:
            printbox_login(session, config)
            parts = urlparse(config["login_url"]); root = f"{parts.scheme}://{parts.netloc}"
            links = {}
            for page in range(1, 20):
                response = session.get(f"{root}/admin/devices?order=id_desc&page={page}&scope=kiosks", timeout=25); response.raise_for_status()
                rows = BeautifulSoup(response.text, "html.parser").select("tbody tr")
                if not rows: break
                for row in rows:
                    cells = row.find_all("td")
                    if len(cells) >= 2:
                        link = cells[1].find("a", href=True)
                        if link:
                            label = link.get_text(" ", strip=True).split(" - ", 1)[0]
                            if label in hosts: links[label] = link["href"]
                if len(rows) < 30: break
            locations = {}
            for host, href in links.items():
                response = session.get(root + href, timeout=25); response.raise_for_status()
                locations[host] = location_fields(BeautifulSoup(response.text, "html.parser"))
        for host in hosts:
            locations.setdefault(host, {"area": "rest", "address": "", "lat": None, "lon": None})
        with cache_lock: location_cache.update(time=time.monotonic(), locations=locations)
        return {host: locations[host] for host in hosts}
    except (ValueError, KeyError, json.JSONDecodeError, requests.RequestException, DashboardError):
        return {host: {"area": "rest", "address": "", "lat": None, "lon": None} for host in hosts}

def normalize_label(value):
    return re.sub(r"[^A-Z0-9]", "", value.upper())


def labels_match(left, right):
    a, b = normalize_label(left), normalize_label(right)
    if a == b: return True
    if abs(len(a) - len(b)) != 1: return False
    short, long = (a, b) if len(a) < len(b) else (b, a)
    return any(long[:index] + long[index + 1:] == short for index in range(len(long)))

def sheet_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def select_terminal_info(rows, hosts):
    result = {}
    for host in hosts:
        match = rows.get(normalize_label(host))
        if not match:
            match = next((row for row in rows.values() if labels_match(host, row.get("Label", ""))), None)
        if match:
            result[host] = match
    return result


def fetch_terminal_info(hosts, force=False):
    now = time.monotonic()
    with cache_lock:
        if not force and sheet_cache["time"] and now - sheet_cache["time"] < 600:
            rows = sheet_cache["rows"]
            return select_terminal_info(rows, hosts)
    try:
        response = requests.get(SHEET_XLSX_URL, timeout=30)
        response.raise_for_status()
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        sheet = workbook["Terminals"] if "Terminals" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [sheet_text(value) for value in next(values)]
        rows = {}
        for source_values in values:
            row = {
                header: sheet_text(value)
                for header, value in zip(headers, source_values)
                if header and sheet_text(value)
            }
            label = row.get("Label", "")
            if label:
                rows[normalize_label(label)] = row
        workbook.close()
        with cache_lock:
            sheet_cache.update(time=time.monotonic(), rows=rows)
        return select_terminal_info(rows, hosts)
    except (StopIteration, ValueError, KeyError, BadZipFile, InvalidFileException, requests.RequestException):
        return {}

def report_locations(html, tbody_index):
    bodies = BeautifulSoup(html, "html.parser").find_all("tbody")
    if len(bodies) <= tbody_index: return set()
    return {row.find_all("td")[0].get_text(strip=True) for row in bodies[tbody_index].find_all("tr") if row.find_all("td")}


def activity_date(html, printing=False):
    table = BeautifulSoup(html, "html.parser").find("table")
    if not table: return None
    headers = [cell.get_text(" ", strip=True) for cell in table.find_all("th")]
    try:
        created_index = headers.index("Created At")
        state_index = headers.index("Final State") if printing else None
    except ValueError: return None
    for row in table.find_all("tr")[1:]:
        cells = row.find_all("td")
        if len(cells) <= created_index: continue
        if state_index is not None and (len(cells) <= state_index or cells[state_index].get_text(strip=True).lower() != "completed"): continue
        try: return datetime.strptime(cells[created_index].get_text(" ", strip=True), "%B %d, %Y %H:%M")
        except ValueError: continue
    return None


def fetch_hidden_problems(zabbix_hosts, force=False):
    now_mono = time.monotonic()
    with cache_lock:
        if not force and hidden_cache["problems"] and now_mono - hidden_cache["time"] < 1800:
            return hidden_cache["problems"]
    try:
        config = json.loads(os.environ.get("PRINTBOX_CONFIG_JSON", ""))
        with requests.Session() as session:
            printbox_login(session, config)
            parts = urlparse(config["login_url"]); root = f"{parts.scheme}://{parts.netloc}"
            devices = {}
            for page in range(1, 20):
                response = session.get(f"{root}/admin/devices?order=id_desc&page={page}&scope=kiosks", timeout=25); response.raise_for_status()
                rows = BeautifulSoup(response.text, "html.parser").select("tbody tr")
                if not rows: break
                for row in rows:
                    cells = row.find_all("td")
                    if len(cells) < 6: continue
                    link = cells[1].find("a", href=True); network = cells[4].find("a", href=True)
                    if not link or not network or not network["href"].rstrip("/").endswith("/networks/14") or cells[5].get_text(strip=True).lower() != "production": continue
                    match = re.search(r"/admin/devices/(\d+)", link["href"])
                    label = link.get_text(" ", strip=True).split(" - ", 1)[0]
                    if match and label: devices[label] = {"id": match.group(1), "deployment": link["href"]}
                if len(rows) < 30: break
            start = date.today() - timedelta(days=10); end = date.today() + timedelta(days=1)
            revenue_url = config["base_target_url"].format(date=start.isoformat(), next_date=end.isoformat())
            printing_url = config["base_printing_url"].format(date=start.isoformat(), next_date=end.isoformat())
            revenue_active = report_locations(session.get(revenue_url, timeout=25).text, 1)
            printing_active = report_locations(session.get(printing_url, timeout=25).text, 0)
            candidates = {label:data for label,data in devices.items() if label not in revenue_active and label not in printing_active and not any(labels_match(label, host) for host in zabbix_hosts)}
            cookies, headers = session.cookies.get_dict(), dict(session.headers)

        def inspect(label, device):
            fill_response = requests.get(f"{root}/admin/devices/{device['id']}/credits_fillups?order=created_at_desc", cookies=cookies, headers=headers, timeout=25); fill_response.raise_for_status()
            print_response = requests.get(f"{root}/admin/devices/{device['id']}/print_logs?order=created_at_desc", cookies=cookies, headers=headers, timeout=25); print_response.raise_for_status()
            fillup, printed = activity_date(fill_response.text), activity_date(print_response.text, True)
            latest = max([value for value in (fillup, printed) if value], default=None)
            days = (datetime.now() - latest).total_seconds() / 86400 if latest else 9999
            if days < 5: return None
            deployment_response = requests.get(root + device["deployment"], cookies=cookies, headers=headers, timeout=25); deployment_response.raise_for_status()
            location = location_fields(BeautifulSoup(deployment_response.text, "html.parser"))
            address, lat, lon = location["address"], location["lat"], location["lon"]
            priority = 4 if days >= 10 else 2
            since = int(latest.timestamp()) if latest else 0
            return {"id": f"hidden-{device['id']}", "host": label, "hostId": device["id"], "description": f"Keine Drucke oder Einzahlungen seit {int(days)} Tagen", "priority": priority, "since": since, "category": "hidden", "area": classify_area(address), "address": address, "lat": lat, "lon": lon, "hidden": True, "lastPrint": printed.isoformat() if printed else None, "lastFillup": fillup.isoformat() if fillup else None}

        hidden = []
        with ThreadPoolExecutor(max_workers=8) as executor:
            futures = [executor.submit(inspect, label, device) for label, device in candidates.items()]
            for future in as_completed(futures):
                try:
                    value = future.result()
                    if value: hidden.append(value)
                except requests.RequestException: continue
        with cache_lock: hidden_cache.update(time=time.monotonic(), problems=hidden)
        return hidden
    except (ValueError, KeyError, json.JSONDecodeError, requests.RequestException, DashboardError):
        return []


def fetch_problems(force=False):
    now = time.monotonic()
    with cache_lock:
        if not force and problem_cache["data"] and now - problem_cache["time"] < CACHE_SECONDS:
            return problem_cache["data"]
    username, password = os.environ.get("ZABBIX_USERNAME"), os.environ.get("ZABBIX_PASSWORD")
    if not username or not password:
        raise DashboardError("Zabbix-Zugangsdaten fehlen.")
    token = zabbix_call("user.login", {"user": username, "password": password})
    triggers = zabbix_call("trigger.get", {
        "output": ["triggerid", "description", "priority", "lastchange"],
        "filter": {"value": 1}, "active": True, "skipDependent": True,
        "monitored": True, "groupids": [os.environ.get("ZABBIX_GROUP_ID", "19")],
        "selectHosts": ["hostid", "host", "name"], "limit": 500,
    }, token, 2)
    problems = []
    for trigger in triggers:
        host_data = (trigger.get("hosts") or [{}])[0]
        host = host_data.get("name") or host_data.get("host") or "Unbekannt"
        description = clean_description(trigger["description"], host)
        problems.append({
            "id": trigger["triggerid"], "host": host, "hostId": host_data.get("hostid"),
            "description": description, "priority": int(trigger["priority"]),
            "since": int(trigger["lastchange"]), "category": category(description),
        })
    locations = fetch_locations({problem["host"] for problem in problems})
    for problem in problems:
        location = locations.get(problem["host"], {})
        problem.update(
            area=location.get("area", "rest"), address=location.get("address", ""),
            lat=location.get("lat"), lon=location.get("lon"),
        )
    problems.extend(fetch_hidden_problems({problem["host"] for problem in problems}, force))
    terminal_info = fetch_terminal_info({problem["host"] for problem in problems}, force)
    for problem in problems:
        problem["terminalInfo"] = terminal_info.get(problem["host"])
    problems.sort(key=lambda x: (-x["priority"], x["since"]))
    result = {"problems": problems, "updatedAt": int(time.time())}
    with cache_lock:
        problem_cache.update(time=time.monotonic(), data=result)
    return result


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(APP_DIR), **kwargs)

    def guess_type(self, path):
        value = super().guess_type(path)
        return f"{value}; charset=utf-8" if value in {"text/html", "text/css", "text/javascript"} else value

    def end_headers(self):
        path = urlparse(self.path).path
        if path == "/" or path.endswith((".html", ".css", ".js")):
            self.send_header("Cache-Control", "no-cache, must-revalidate")
        super().end_headers()

    def send_json(self, payload, status=HTTPStatus.OK, cookie=None):
        body = json.dumps(payload, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        if cookie:
            self.send_header("Set-Cookie", cookie)
        self.end_headers()
        self.wfile.write(body)

    def read_json(self):
        length = int(self.headers.get("Content-Length", 0))
        if length > 10000:
            raise ValueError("Anfrage zu groß.")
        return json.loads(self.rfile.read(length))

    def current_user(self):
        secret = os.environ.get("SESSION_SECRET", "")
        cookie = self.headers.get("Cookie", "")
        token = next((x.split("=", 1)[1] for x in cookie.split("; ") if x.startswith("zbx_session=")), "")
        try:
            raw, signature = token.rsplit(".", 1)
            if not hmac.compare_digest(signature, hmac.new(secret.encode(), raw.encode(), hashlib.sha256).hexdigest()):
                return None
            user, expires = base64.urlsafe_b64decode(raw + "==").decode().split("|", 1)
            return user if int(expires) > time.time() else None
        except (ValueError, UnicodeError):
            return None

    def check_password(self, user, password):
        try:
            stored = json.loads(os.environ.get("APP_USERS_JSON", "{}"))[user]
            salt, expected = stored.split("$", 1)
            actual = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 600000).hex()
            return hmac.compare_digest(actual, expected)
        except (KeyError, ValueError, TypeError, json.JSONDecodeError):
            return False

    def do_GET(self):
        path = urlparse(self.path).path
        if path == "/health":
            self.send_json({"status": "ok", "version": os.environ.get("RENDER_GIT_COMMIT", "local")[:7]})
        elif path == "/api/session":
            user = self.current_user(); self.send_json({"authenticated": bool(user), "user": user})
        elif path == "/api/problems":
            if not self.current_user():
                self.send_json({"error": "Nicht angemeldet."}, HTTPStatus.UNAUTHORIZED); return
            try:
                self.send_json(fetch_problems("refresh=1" in self.path))
            except (DashboardError, requests.RequestException) as exc:
                self.send_json({"error": str(exc)}, HTTPStatus.BAD_GATEWAY)
        else:
            super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path
        if path == "/api/login":
            try:
                data = self.read_json(); user = str(data.get("username", "")).strip()
                if not self.check_password(user, str(data.get("password", ""))):
                    time.sleep(.7); self.send_json({"error": "Falsche Anmeldung."}, HTTPStatus.UNAUTHORIZED); return
                expires = int(time.time()) + 28800
                raw = base64.urlsafe_b64encode(f"{user}|{expires}".encode()).decode().rstrip("=")
                signature = hmac.new(os.environ["SESSION_SECRET"].encode(), raw.encode(), hashlib.sha256).hexdigest()
                secure = "; Secure" if os.environ.get("RENDER") else ""
                self.send_json({"authenticated": True}, cookie=f"zbx_session={raw}.{signature}; Path=/; HttpOnly; SameSite=Lax; Max-Age=28800{secure}")
            except (ValueError, json.JSONDecodeError):
                self.send_json({"error": "Ungültige Anfrage."}, HTTPStatus.BAD_REQUEST)
        elif path == "/api/logout":
            self.send_json({}, cookie="zbx_session=; Path=/; HttpOnly; SameSite=Lax; Max-Age=0")
        else:
            self.send_error(HTTPStatus.NOT_FOUND)


def main():
    for key in ("SESSION_SECRET", "APP_USERS_JSON", "ZABBIX_USERNAME", "ZABBIX_PASSWORD"):
        if not os.environ.get(key):
            raise SystemExit(f"{key} fehlt.")
    host = "0.0.0.0" if os.environ.get("RENDER") else "127.0.0.1"
    server = ThreadingHTTPServer((host, int(os.environ.get("PORT", "8081"))), Handler)
    print("Zabbix Dashboard läuft.")
    server.serve_forever()


if __name__ == "__main__":
    main()
